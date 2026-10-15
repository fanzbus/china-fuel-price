#!/usr/bin/env python3
"""每周抓一次全球主要国家油价：EU Weekly Oil Bulletin + FRED GASREGW + 本站中国均价。

规则（诚实性硬要求）：
- 只取实测到的数据，不编造。EU 只有 95 号 + 柴油，不用 95 冒充 92/98；美国只有汽油，无柴油。
- 日本/韩国/印度/巴西等暂无免费稳定官方零售价源，保持待接入，只外链，不搬数据。
- 任一上游失败 -> 该块缺失，status=partial，记入 errors；EU 与 FRED 都失败 -> 不覆盖旧文件，直接失败退出。
- 汇率失败 -> 无法换算，不覆盖旧文件，直接失败退出。
- 每周跑一次（GitHub Actions 周一定时），不要高频。
"""
import json
import re
import sys
import urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

CST = timezone(timedelta(hours=8))
ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
PUBLIC_DATA = ROOT / "public" / "data"
UA = "china-fuel-price/1.0 (+https://github.com/fanzbus/china-fuel-price; weekly global sync, contact via repo)"

EU_PAGE = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
FRED_GASREGW = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=GASREGW"
ER_API = "https://open.er-api.com/v6/latest/USD"
GAL_TO_LITER = 3.78541

EU_CN = {
    "Austria": "奥地利", "Belgium": "比利时", "Bulgaria": "保加利亚", "Croatia": "克罗地亚",
    "Cyprus": "塞浦路斯", "Czechia": "捷克", "Denmark": "丹麦", "Estonia": "爱沙尼亚",
    "Finland": "芬兰", "France": "法国", "Germany": "德国", "Greece": "希腊",
    "Hungary": "匈牙利", "Ireland": "爱尔兰", "Italy": "意大利", "Latvia": "拉脱维亚",
    "Lithuania": "立陶宛", "Luxembourg": "卢森堡", "Malta": "马耳他", "Netherlands": "荷兰",
    "Poland": "波兰", "Portugal": "葡萄牙", "Romania": "罗马尼亚", "Slovakia": "斯洛伐克",
    "Slovenia": "斯洛文尼亚", "Spain": "西班牙", "Sweden": "瑞典",
}
PENDING_COUNTRIES = ["日本", "韩国", "印度", "巴西"]


def fetch_text(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read().decode("utf-8", errors="replace")


def fetch_bytes(url, timeout=60):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def fetch_eu(rates_cny, rates_eur):
    """欧盟 Weekly Oil Bulletin：先抓列表页解析最新含税 xlsx 链接（UUID 每周变），再下载解析。

    xlsx 结构：行1表头（B 列 Euro-super 95，C 列 Diesel），行2 A 列为数据日期，
    行3起各国，B/C 列为 EUR/1000 升。跳过 CE/EC 开头的加权平均汇总行。
    """
    import openpyxl

    page = fetch_text(EU_PAGE)
    m = re.search(
        r"document/download/([0-9a-f\-]+_en\?filename=Weekly%20Oil%20Bulletin%20Weekly%20prices%20with%20Taxes%20-[^\"&\s]*\.xlsx)",
        page)
    if not m:
        raise RuntimeError("EU bulletin xlsx link not found on list page")
    xlsx_url = "https://energy.ec.europa.eu/document/download/" + m.group(1)
    raw = fetch_bytes(xlsx_url)
    wb = openpyxl.load_workbook(io_bytes(raw), data_only=True)
    ws = wb.active
    date_cell = ws.cell(row=2, column=1).value
    if hasattr(date_cell, "strftime"):
        as_of = date_cell.strftime("%Y-%m-%d")
    else:
        as_of = str(date_cell)[:10]
    eur_per_usd = rates_eur  # ER-API: EUR per 1 USD
    cny_per_eur = rates_cny / rates_eur
    rows = []
    for r in range(3, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or not isinstance(name, str):
            continue
        if name.startswith("CE/EC") or "Moyenne" in name or "Weighted" in name:
            continue  # 加权平均汇总行，不是国家
        if name not in EU_CN:
            continue
        v95 = ws.cell(row=r, column=2).value
        vd = ws.cell(row=r, column=3).value
        if v95 is None and vd is None:
            continue
        def conv(v):
            if v is None:
                return None, None
            eur_l = float(v) / 1000.0
            return round(eur_l * cny_per_eur, 2), round(eur_l / eur_per_usd, 3)
        p95_cny, p95_usd = conv(v95)
        d_cny, d_usd = conv(vd)
        rows.append({"country": name, "country_cn": EU_CN[name],
                     "p95_cny": p95_cny, "p95_usd": p95_usd,
                     "diesel_cny": d_cny, "diesel_usd": d_usd,
                     "unit": "元/升", "as_of": as_of,
                     "source": "欧盟委员会 Weekly Oil Bulletin（含税价）",
                     "note": "Euro-super 95"})
    if not rows:
        raise RuntimeError("EU xlsx parsed 0 countries")
    return rows, as_of


def io_bytes(raw):
    import io
    return io.BytesIO(raw)


def fetch_us(rates_cny):
    """FRED GASREGW：美国零售汽油（周），USD/加仑 -> USD/升 -> CNY/升。"""
    csv_text = fetch_text(FRED_GASREGW)
    last = None
    for line in csv_text.strip().splitlines()[1:]:
        parts = line.split(",")
        if len(parts) < 2:
            continue
        d, v = parts[0].strip(), parts[1].strip()
        if v in ("", "."):
            continue
        try:
            last = (d, float(v))
        except ValueError:
            continue
    if not last:
        raise RuntimeError("FRED GASREGW no usable rows")
    date, usd_gal = last
    usd_l = usd_gal / GAL_TO_LITER
    return {"country": "United States", "country_cn": "美国",
            "p95_cny": round(usd_l * rates_cny, 2), "p95_usd": round(usd_l, 3),
            "diesel_cny": None, "diesel_usd": None,
            "unit": "元/升", "as_of": date,
            "source": "FRED GASREGW（美联储）",
            "note": "常规汽油（约92号），无柴油数据"}


def fetch_china(rates_cny):
    """本站中国全国均价（qiyoujiage.com 31 省），取自 data/prices.json。"""
    prices = json.loads((DATA / "prices.json").read_text(encoding="utf-8"))
    provs = prices["provinces"]
    avg = lambda k: round(sum(r[k] for r in provs) / len(provs), 2)
    as_of = (prices.get("fetched_at") or "")[:10]
    return {"country": "China", "country_cn": "中国",
            "p92_cny": avg("p92"), "p95_cny": avg("p95"), "p98_cny": avg("p98"),
            "diesel_cny": avg("d0"),
            "p92_usd": round(avg("p92") / rates_cny, 3),
            "p95_usd": round(avg("p95") / rates_cny, 3),
            "p98_usd": round(avg("p98") / rates_cny, 3),
            "diesel_usd": round(avg("d0") / rates_cny, 3),
            "unit": "元/升", "as_of": as_of,
            "source": "本站全国均价（qiyoujiage.com 31 省）",
            "note": "92/95/98/柴油均为全国均价"}


def main():
    fetched_at = datetime.now(CST).isoformat(timespec="seconds")
    try:
        fx = json.loads(fetch_text(ER_API, timeout=15))
        rates = fx.get("rates", {})
        rates_cny, rates_eur = rates.get("CNY"), rates.get("EUR")
        if not rates_cny or not rates_eur:
            raise RuntimeError("ER-API missing CNY/EUR rates")
    except Exception as e:  # noqa: BLE001
        print(f"FAILED: FX unavailable, keep old global.json: {e}", file=sys.stderr)
        return 2

    countries, errors = [], []
    eu_as_of = None
    try:
        eu_rows, eu_as_of = fetch_eu(rates_cny, rates_eur)
        countries.extend(eu_rows)
        print(f"EU OK: {len(eu_rows)} countries as_of={eu_as_of}")
    except Exception as e:  # noqa: BLE001
        errors.append(f"EU Bulletin 失败: {e}")
        print(f"EU failed: {e}")
    try:
        countries.append(fetch_us(rates_cny))
        print("US OK (FRED GASREGW)")
    except Exception as e:  # noqa: BLE001
        errors.append(f"FRED GASREGW 失败: {e}")
        print(f"US failed: {e}")
    try:
        china = fetch_china(rates_cny)
        countries.insert(0, china)
        print("China OK (本站全国均价）")
    except Exception as e:  # noqa: BLE001
        errors.append(f"中国均价失败: {e}")
        print(f"China failed: {e}")

    if not countries or (not eu_as_of and not any(c["country"] == "United States" for c in countries)):
        # EU 与 FRED 都失败 -> 不覆盖旧文件
        print(f"FAILED: EU and FRED both unavailable, keep old global.json. errors={errors}", file=sys.stderr)
        return 2
    # 排序：中国、美国在前，其余按中文名
    order = {"中国": 0, "美国": 1}
    countries.sort(key=lambda c: (order.get(c["country_cn"], 2), c["country_cn"]))
    out = {"fetched_at": fetched_at, "status": "ok" if not errors else "partial",
           "as_of": eu_as_of or (countries[0].get("as_of") if countries else None),
           "fx_usd_cny": rates_cny, "fx_eur_per_usd": rates_eur,
           "sources": ["欧盟委员会 Weekly Oil Bulletin（含税价，95 号 + 柴油）",
                       "FRED GASREGW（美国常规汽油）",
                       "本站全国均价（qiyoujiage.com 31 省）"],
           "countries": countries,
           "pending": PENDING_COUNTRIES,
           "pending_note": "日本/韩国/印度/巴西等暂无免费稳定官方零售价源，保持待接入，不编造数据。可在 GlobalPetrolPrices 查看。",
           "errors": errors}
    (DATA / "global.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    (PUBLIC_DATA / "global.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    # meta：global_status + 说明（保留其余字段）
    meta_path = DATA / "meta.json"
    meta = {}
    if meta_path.exists():
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            meta = {}
    meta["global_status"] = "ok" if not errors else "partial"
    meta.setdefault("data_quality", {})["global"] = (
        f"全球对比已接入：欧盟 27 国（Weekly Oil Bulletin）+ 美国（FRED）+ 中国（本站均价）；"
        f"数据日期见各行 as_of；{('部分来源失败：' + '；'.join(errors)) if errors else ''}")
    (meta_path).write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (PUBLIC_DATA / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK countries={len(countries)} status={out['status']} errors={len(errors)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
